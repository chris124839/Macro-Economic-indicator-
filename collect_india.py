#!/usr/bin/env python3
"""Populate India's CPI, WPI inflation and RBI policy-rate history from RBI DBIE.

The importer fetches data at run time in GitHub Actions. It selects CPI Combined / All
India General Index and WPI All Commodities only; it does not synthesize observations.
Existing remarks are preserved when official records are refreshed.
"""
from __future__ import annotations
import calendar, csv, io, json, re, sys, time
from datetime import date, datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen
from zipfile import BadZipFile
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "data" / "observations.csv"
API = "https://data-api.dbie.rbihub.in"
CPI_WORKBOOK = ("https://raw.githubusercontent.com/Reserve-Bank-Innovation-Hub/"
                "dbie.rbihub.in/dev/data/publications/monthly-rbi-bulletin/"
                "prices-and-production/table-19-consumer-price-index-base-2010-100.xlsx")
URLS = {
    "rbi_repo": "https://dbie.rbihub.in/tables?table=financial_sector.r531_key_rates",
    "rbi_reverse_repo": "https://dbie.rbihub.in/tables?table=financial_sector.r531_key_rates",
    "rbi_sdf": "https://dbie.rbihub.in/tables?table=financial_sector.r531_key_rates",
    "india_cpi_yoy": "https://dbie.rbihub.in/prices/consumer-price-index",
    "india_wpi_yoy": "https://dbie.rbihub.in/tables/wholesale-price-index-inflation",
}
COLUMNS = ["indicator","country","period_date","release_date","value","unit","source_url","source_name","remarks"]
HEADERS = {"User-Agent":"IndiaMacroDashboard/1.0 (historical data import)"}


def norm(value):
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def download(url, retries=3):
    error = None
    for attempt in range(retries):
        try:
            req = Request(url, headers=HEADERS)
            with urlopen(req, timeout=60) as response:
                data = response.read()
            if not data:
                raise RuntimeError("empty response")
            return data
        except (HTTPError, URLError, TimeoutError, OSError, RuntimeError) as exc:
            error = exc
            if attempt + 1 < retries:
                time.sleep(2 ** attempt)
    raise RuntimeError(f"Download failed for {url}: {error}")


def fetch_csv(url):
    text = download(url).decode("utf-8-sig", errors="replace")
    if text.lstrip().startswith("{"):
        try:
            payload = json.loads(text)
            raise RuntimeError(f"DBIE API returned an error: {payload}")
        except json.JSONDecodeError:
            pass
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise RuntimeError(f"No CSV header returned by {url}")
    return [{str(k or "").strip(): (v or "").strip() for k, v in row.items() if k is not None}
            for row in reader]


def find_column(row, candidates):
    mapping = {norm(k): k for k in row}
    for candidate in candidates:
        if norm(candidate) in mapping:
            return mapping[norm(candidate)]
    for candidate in candidates:
        key = norm(candidate)
        for nk, original in mapping.items():
            if nk.endswith("_" + key) or nk.startswith(key + "_"):
                return original
    return None


def parse_date(value):
    s = str(value or "").strip()
    if not s or s.lower() in {"-", "—", "na", "n/a", "none", "null"}:
        return None
    for fmt in ("%Y-%m-%d", "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y", "%d/%m/%Y",
                "%d-%b-%Y", "%d %b %Y", "%d-%B-%Y", "%b-%Y", "%B-%Y",
                "%b %Y", "%B %Y", "%Y-%m"):
        try:
            return datetime.strptime(s[:19], fmt).date()
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s.replace("Z", "+00:00")).date()
    except ValueError:
        return None


def parse_num(value):
    s = str(value or "").strip().replace(",", "").replace("%", "")
    if not s or s.lower() in {"-", "—", "na", "n/a", "none", "null"}:
        return None
    try:
        n = float(s)
        return n if n == n and abs(n) != float("inf") else None
    except ValueError:
        return None


def observation(indicator, period, value, unit, source_name):
    return {"indicator": indicator, "country": "India", "period_date": period,
            "release_date": "", "value": value, "unit": unit,
            "source_url": URLS[indicator], "source_name": source_name, "remarks": ""}


def fetch_rates():
    """Read RBI DBIE Key Rates from typed columns or its spreadsheet-grid export.

    DBIE can expose this particular report as c1..c8 rather than named rate columns.
    The published Key Rates table's leaf-column order is:
      c1 Effective Date, c2 Bank Rate, c3 Repo, c4 Reverse Repo,
      c5 SDF, c6 MSF, c7 CRR, c8 SLR.
    The parent heading "Fix Range LAF Rates" spans the Repo and Reverse Repo columns.
    A dash in a rate cell means no change, so the last published value is carried
    forward. SDF remains blank until its first actual value in April 2022.
    """
    rows = fetch_csv(f"{API}/api/tables/financial_sector/r531_key_rates/csv")
    if not rows:
        raise RuntimeError("RBI Key Rates table returned no rows")

    first = rows[0]
    date_col = find_column(first, ["effective_date", "effective date", "date", "time_period", "period"])
    rate_cols = {
        "rbi_repo": find_column(first, ["repo_rate", "repo", "repo rate", "policy repo rate"]),
        "rbi_reverse_repo": find_column(first, ["reverse_repo_rate", "reverse_repo", "reverse repo", "reverse repo rate"]),
        "rbi_sdf": find_column(first, ["standing_deposit_facility_sdf_rate", "standing deposit facility sdf rate", "sdf rate"]),
    }

    # Case 1: typed/normalized API response with useful column names.
    if date_col and all(rate_cols.values()):
        dated = [(d, row) for row in rows if (d := parse_date(row.get(date_col))) is not None]
        dated.sort(key=lambda item: item[0])
        if not dated:
            raise RuntimeError("No effective dates parsed from RBI Key Rates table")
        current, result = {key: None for key in rate_cols}, []
        for dt, row in dated:
            for indicator, col in rate_cols.items():
                value = parse_num(row.get(col))
                if value is not None:
                    current[indicator] = value
                if current[indicator] is not None:
                    result.append(observation(indicator, dt.isoformat(), current[indicator], "%",
                                              "Reserve Bank of India (DBIE) — Key Rates"))
        counts = {k: sum(r["indicator"] == k for r in result) for k in rate_cols}
        missing = [k for k, count in counts.items() if count == 0]
        if missing:
            raise RuntimeError(f"Typed Key Rates data missing series: {missing}; counts={counts}")
        print(f"RBI Key Rates (typed): {len(dated)} dates; counts={counts}; earliest={dated[0][0]}")
        return result

    # Case 2: faithful spreadsheet-grid export (src_file, tab, period, row_no, c1..c8).
    # Do not infer columns from merged group headings: their text is stored in different
    # rows/cells and is the reason the previous parser failed. Use the official leaf
    # column positions and validate that date and expected rates exist before importing.
    grid_cols = sorted(
        [col for col in first if re.fullmatch(r"c\d+", norm(col))],
        key=lambda col: int(re.search(r"\d+", col).group())
    )
    if len(grid_cols) < 5:
        raise RuntimeError(f"Key Rates report-grid expected c1..c8; headers={list(first)}")

    # Map by numeric suffix (so it remains correct if the export's dict order changes).
    by_number = {int(re.search(r"\d+", col).group()): col for col in grid_cols}
    if not all(n in by_number for n in (1, 3, 4, 5)):
        raise RuntimeError(f"Key Rates report-grid lacks c1/date, c3/repo, c4/reverse repo or c5/SDF; columns={grid_cols}")

    # There is usually one sheet group. Preserve separate sheet/file groups to guard
    # against future DBIE table layout changes and deduplicate after parsing.
    groups = {}
    for row in rows:
        key = (str(row.get("src_file", "")).strip(), str(row.get("tab", "")).strip())
        groups.setdefault(key, []).append(row)

    result = []
    parsed_groups = 0
    for (source_file, tab_name), sheet_rows in groups.items():
        def row_number(item):
            try:
                return int(float(item.get("row_no", 0)))
            except (TypeError, ValueError):
                return 0

        sheet_rows.sort(key=row_number)
        dated = []
        for row in sheet_rows:
            # In RBI Key Rates workbook c1 is the effective-date column. Header rows,
            # footnotes and titles fail date parsing and are naturally skipped.
            dt = parse_date(row.get(by_number[1]))
            if dt:
                dated.append((dt, row))
        if not dated:
            continue
        dated.sort(key=lambda item: item[0])

        # Official DBIE Key Rates leaf-column positions (as listed on the RBI page):
        # c2 Bank Rate, c3 Repo, c4 Reverse Repo, c5 SDF, c6 MSF, c7 CRR, c8 SLR.
        # The requested dashboard uses repo, reverse repo and SDF only.
        col_map = {
            "rbi_repo": by_number.get(3),
            "rbi_reverse_repo": by_number.get(4),
            "rbi_sdf": by_number.get(5),
        }
        current = {key: None for key in col_map}
        for dt, row in dated:
            for indicator, col in col_map.items():
                if not col:
                    continue
                value = parse_num(row.get(col))
                # Per the RBI table notes, '-' means unchanged. Carry forward the last
                # rate change. For SDF, this starts only on the first row with a value.
                if value is not None:
                    current[indicator] = value
                if current[indicator] is not None:
                    result.append(observation(
                        indicator, dt.isoformat(), current[indicator], "%",
                        "Reserve Bank of India (DBIE) — Key Rates"
                    ))
        parsed_groups += 1

    dedup = {}
    for row in result:
        dedup[(row["indicator"], row["period_date"])] = row
    result = sorted(dedup.values(), key=lambda row: (row["indicator"], row["period_date"]))
    counts = {key: sum(row["indicator"] == key for row in result)
              for key in ("rbi_repo", "rbi_reverse_repo", "rbi_sdf")}
    missing = [key for key, count in counts.items() if count == 0]
    if missing:
        # Provide a compact but useful diagnostic without dumping hundreds of rows.
        sample = [{k: row.get(k, "") for k in ("src_file", "tab", "period", "row_no", *grid_cols)}
                  for row in rows[:12]]
        raise RuntimeError(
            f"Key Rates report-grid parser found {parsed_groups} sheets with dates, but missing {missing}; "
            f"headers={list(first)}; sample first rows={sample}"
        )
    periods = [row["period_date"] for row in result]
    print(f"RBI Key Rates (report grid): sheets={parsed_groups}; counts={counts}; "
          f"dates={min(periods)} through {max(periods)}")
    return result


def fetch_wpi():
    # All Commodities = WPI01, current standard base = 2011-12.
    params = urlencode({"base_per":"BY_2011_12", "wpi_com_cls_rn":"WPI01"})
    rows = fetch_csv(f"{API}/api/tables/real_sector/whole_price_index_inf_rn/csv?{params}")
    if not rows:
        raise RuntimeError("WPI API returned no All Commodities rows for base 2011-12")
    first = rows[0]
    date_col = find_column(first, ["time_period", "period", "date", "observation_date"])
    value_col = find_column(first, ["obs_value", "value", "inflation", "rate"])
    if not date_col or not value_col:
        raise RuntimeError(f"WPI columns not recognized; headers={list(first)}")
    base_col = find_column(first, ["base_per"])
    class_col = find_column(first, ["wpi_com_cls_rn"])
    base_label_col = find_column(first, ["base_per_label", "base_per_description"])
    class_label_col = find_column(first, ["wpi_com_cls_rn_label", "wpi_com_cls_rn_description"])
    result = []
    for row in rows:
        base = norm(row.get(base_col, "")) if base_col else ""
        base_label = norm(row.get(base_label_col, "")) if base_label_col else ""
        cls = norm(row.get(class_col, "")) if class_col else ""
        cls_label = norm(row.get(class_label_col, "")) if class_label_col else ""
        if base and base not in {"by_2011_12", "base_year_2011_12"} and "2011_12" not in base and "2011_12" not in base_label:
            continue
        if cls and cls not in {"wpi01", "all_commodities"} and "all_commodities" not in cls_label:
            continue
        dt, val = parse_date(row.get(date_col)), parse_num(row.get(value_col))
        if dt and val is not None:
            result.append(observation("india_wpi_yoy", dt.isoformat(), val, "% YoY",
                                      "Reserve Bank of India (DBIE) — WPI Inflation, All Commodities, base 2011-12"))
    by_date = {r["period_date"]: r for r in result}
    result = [by_date[k] for k in sorted(by_date)]
    if not result:
        raise RuntimeError(f"No headline WPI observations parsed; headers={list(first)}")
    print(f"WPI: {len(result)} monthly observations, {result[0]['period_date']} through {result[-1]['period_date']}")
    return result


def parse_month(value):
    m = re.fullmatch(r"([A-Z]{3})[- ](\d{4})", str(value or "").strip().upper())
    if not m:
        return None
    month = {name: i for i, name in enumerate(
        ["JAN","FEB","MAR","APR","MAY","JUN","JUL","AUG","SEP","OCT","NOV","DEC"], 1)}.get(m.group(1))
    return f"{int(m.group(2)):04d}-{month:02d}" if month else None


def end_of_month(yyyy_mm):
    y, m = map(int, yyyy_mm.split("-"))
    return date(y, m, calendar.monthrange(y, m)[1]).isoformat()


def fetch_cpi():
    raw = download(CPI_WORKBOOK)
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except (BadZipFile, ValueError) as exc:
        raise RuntimeError(f"CPI workbook cannot be parsed: {exc}") from exc
    choices = {}
    # CPI Table 19 columns: B=month, C=commodity, D=status, J=Combined inflation.
    for ws in wb.worksheets:
        m = re.match(r"^CPI\s*-\s*(\d{4})=(\d+)\s*\(All India\)$", ws.title.strip(), re.I)
        if not m:
            continue
        base_year = int(m.group(1))
        base_label = f"{m.group(1)}={m.group(2)}"
        for values in ws.iter_rows(min_row=8, values_only=True):
            if len(values) < 10:
                continue
            month = parse_month(values[1])
            commodity = re.sub(r"\s+", " ", str(values[2] or "").strip()).lower()
            if not month or commodity != "a) general index":
                continue
            value = parse_num(values[9])
            if value is None:
                continue
            status = str(values[3] or "").strip().upper()
            choices.setdefault(month, []).append((base_year, value, base_label, status))
    wb.close()
    if not choices:
        raise RuntimeError("No All India General Index / Combined CPI inflation values in RBI DBIE Table 19 workbook")
    result = []
    for period in sorted(choices):
        # Prefer the newer base year for overlapping months; prefer Final over Provisional
        # within a base. Missing early YoY values in a newer base aren't imported as zero.
        base_year, value, base_label, status = sorted(choices[period], key=lambda x: (x[0], x[3] == "F"), reverse=True)[0]
        result.append(observation("india_cpi_yoy", end_of_month(period), value, "% YoY",
                                  f"Reserve Bank of India (DBIE) — CPI Table 19, base {base_label}"))
    print(f"CPI: {len(result)} monthly observations, {result[0]['period_date']} through {result[-1]['period_date']}")
    return result


def merge_save(new_rows):
    existing = []
    if OUT.exists():
        with OUT.open("r", encoding="utf-8-sig", newline="") as f:
            existing = list(csv.DictReader(f))
    merged = {}
    for row in existing:
        key = (str(row.get("indicator", "")).strip(), str(row.get("period_date", "")).strip())
        if all(key):
            merged[key] = {c: row.get(c, "") for c in COLUMNS}
    for row in new_rows:
        key = (row["indicator"], row["period_date"])
        prior = merged.get(key, {})
        clean = {c: row.get(c, "") for c in COLUMNS}
        if not str(clean.get("remarks", "")).strip():
            clean["remarks"] = prior.get("remarks", "")
        merged[key] = clean
    ordered = sorted(merged.values(), key=lambda r: (r["indicator"], r["period_date"]))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    with OUT.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=COLUMNS)
        writer.writeheader()
        writer.writerows(ordered)
    counts = {}
    for row in ordered:
        counts[row["indicator"]] = counts.get(row["indicator"], 0) + 1
    print(f"Saved {len(ordered)} unique observations in {OUT.relative_to(ROOT)}")
    print("Data counts: " + json.dumps(counts, sort_keys=True))


def main():
    all_rows, errors = [], []
    for name, fn in [("RBI Key Rates", fetch_rates), ("WPI inflation", fetch_wpi), ("CPI inflation", fetch_cpi)]:
        try:
            rows = fn()
            if not rows:
                raise RuntimeError("Source returned no usable rows")
            all_rows.extend(rows)
        except Exception as exc:
            errors.append(f"{name}: {exc}")
            print(f"ERROR [{name}]: {exc}")
    if not all_rows:
        print("No official records retrieved; leaving the observations file unchanged.")
        sys.exit(1)
    merge_save(all_rows)
    loaded = {r["indicator"] for r in all_rows}
    required = {"india_cpi_yoy", "india_wpi_yoy", "rbi_repo", "rbi_reverse_repo", "rbi_sdf"}
    missing = sorted(required - loaded)
    if errors or missing:
        print("INCOMPLETE IMPORT — workflow will fail visibly; successful records were written for commit.")
        for e in errors:
            print(" - " + e)
        if missing:
            print(" - Missing series: " + ", ".join(missing))
        sys.exit(1)
    print("SUCCESS: CPI, WPI, repo, reverse repo and SDF histories imported from RBI DBIE.")

if __name__ == "__main__":
    main()
